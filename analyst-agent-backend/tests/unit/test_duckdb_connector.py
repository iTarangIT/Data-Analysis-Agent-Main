"""The file connector, and the lockdown that is the first of its two security layers.

The guard's table allowlist is the second. These tests cover the engine layer: after ingest the
connection cannot reach the filesystem at all, however the query is written.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.connectors.duckdb import (
    TABLE_NAME,
    Dataset,
    DuckDBConnector,
    FileSource,
    Part,
    _numbers,
    ingest_upload,
)
from app.services.errors import DomainError

CSV = "region,product,units,revenue_inr\nWest,Cell,10,2500.50\nEast,Pack,4,1800.00\n"


def _csv(tmp_path: Path, filename: str, body: str, taken: set[str] | None = None):
    src = tmp_path / filename
    src.write_text(body, encoding="utf-8")
    return ingest_upload(src, tmp_path / "out", filename, set() if taken is None else taken)


def _workbook(tmp_path: Path, filename: str, sheets: dict[str, list[list]]):
    book = Workbook()
    for i, (title, rows) in enumerate(sheets.items()):
        sheet = book.active if i == 0 else book.create_sheet(title)
        sheet.title = title
        for row in rows:
            sheet.append(row)
    src = tmp_path / filename
    book.save(src)
    return ingest_upload(src, tmp_path / "out", filename, set())


def _dataset(sources: list[FileSource], root: Path | None = None) -> Dataset:
    return Dataset(
        "t_a",
        [
            Part(
                table=s.table,
                file=s.file,
                origin=s.origin,
                storage_key=Path(s.path).relative_to(root).as_posix() if root else s.path,
                sha256="",
                profile=s.profile,
            )
            for s in sources
        ],
    )


def _types(sources: list[FileSource]) -> dict[str, str]:
    (table,) = _dataset(sources).read_tables([sources[0].table])
    return {c.name: c.type for c in table.columns}


@pytest.fixture
def sales(tmp_path) -> list[FileSource]:
    return _csv(tmp_path, "Q3 sales.csv", CSV)


@pytest.fixture
def connector(sales) -> DuckDBConnector:
    return DuckDBConnector("t_a", sales)


@pytest.fixture
def dataset(sales) -> Dataset:
    return _dataset(sales)


class TestIngest:
    def test_a_csv_becomes_one_parquet_table(self, sales, tmp_path):
        assert [s.table for s in sales] == ["q3_sales"]
        assert Path(sales[0].path).exists()
        assert Path(sales[0].path).suffix == ".parquet"

    def test_a_source_records_its_file_and_origin(self, sales):
        assert (sales[0].file, sales[0].origin) == ("Q3 sales.csv", "upload")

    def test_the_table_name_is_derived_not_taken_from_the_filename(self, tmp_path):
        assert _csv(tmp_path, "weird name!! 2026.csv", CSV)[0].table == "weird_name_2026"

    def test_a_name_that_would_not_be_an_identifier_is_prefixed(self, tmp_path):
        assert _csv(tmp_path, "2026 report.csv", CSV)[0].table == "t_2026_report"

    def test_a_semicolon_separated_file_is_split_into_its_columns(self, tmp_path):
        sources = _csv(tmp_path, "stock.csv", "dealer;units\nA;10\nB;4\n")

        assert list(_types(sources)) == ["dealer", "units"]

    def test_a_parquet_upload_is_read_like_any_other(self, tmp_path, sales):
        (source,) = ingest_upload(Path(sales[0].path), tmp_path / "again", "orders.parquet", set())

        assert source.table == "orders"
        assert source.profile["columns"] == sales[0].profile["columns"]

    def test_a_single_sheet_workbook_is_named_after_the_file(self, tmp_path):
        sources = _workbook(
            tmp_path, "Dealers.xlsx", {"Sheet1": [["dealer", "units"], ["A", 3]], "Notes": []}
        )

        assert [s.table for s in sources] == ["dealers"]

    def test_every_sheet_of_a_spreadsheet_becomes_a_table(self, tmp_path):
        """Answering from the first sheet and silently ignoring the rest is the worst available
        failure for a spreadsheet-only customer."""
        sources = _workbook(
            tmp_path,
            "book.xlsx",
            {
                "Sales": [["a"], [1], [2]],
                "Returns 2026": [["b"], [3]],
                "Stock": [["c"], [4]],
            },
        )

        assert sorted(s.table for s in sources) == [
            "book__returns_2026",
            "book__sales",
            "book__stock",
        ]

    def test_a_name_already_in_the_dataset_gets_a_suffix(self, tmp_path):
        taken = {"q3_sales"}

        sources = _csv(tmp_path, "Q3 sales.csv", CSV, taken)

        assert sources[0].table == "q3_sales_2"
        assert taken == {"q3_sales", "q3_sales_2"}

    def test_a_suffix_on_a_long_name_stays_a_valid_identifier(self, tmp_path):
        stem = "dealer_" * 10
        taken = {stem.strip("_")[:63]}

        (source,) = _csv(tmp_path, f"{stem}.csv", CSV, taken)

        assert source.table.endswith("_2")
        assert TABLE_NAME.match(source.table)

    def test_a_keyword_is_never_a_bare_identifier(self, tmp_path):
        sources = _csv(tmp_path, "order.csv", "group,units\nWest,10\nEast,4\n")

        cols, rows = DuckDBConnector("t_a", sources).run_select(
            "select group_, units from order_ order by units", 10
        )

        assert cols == ["group_", "units"]
        assert rows == [("East", 4), ("West", 10)]

    def test_an_empty_file_is_refused_with_a_readable_reason(self, tmp_path):
        import pandas as pd

        src = tmp_path / "empty.xlsx"
        with pd.ExcelWriter(src) as writer:
            pd.DataFrame().to_excel(writer, sheet_name="Blank", index=False)

        with pytest.raises(DomainError, match=r"empty\.xlsx has no readable rows"):
            ingest_upload(src, tmp_path / "out", src.name, set())


class TestHardening:
    def test_a_title_above_the_header_is_skipped(self, tmp_path):
        (source,) = _workbook(
            tmp_path,
            "dealer.xlsx",
            {
                "July": [
                    ["Dealer sales report"],
                    ["Period: July 2026"],
                    ["Dealer", "Revenue (INR)", "Units"],
                    ["Pune Motors", 2500.5, 10],
                    ["Nashik EV", 1800.0, 4],
                ]
            },
        )

        assert source.profile["header_row"] == 2
        assert source.profile["columns"] == {
            "dealer": "Dealer",
            "revenue_inr": "Revenue (INR)",
            "units": "Units",
        }
        assert source.profile["row_count"] == 2

    def test_a_grand_total_row_is_dropped_and_counted(self, tmp_path):
        sources = _workbook(
            tmp_path,
            "dealer.xlsx",
            {"July": [["dealer", "units"], ["A", 10], ["B", 4], ["Grand Total", 14]]},
        )

        _, rows = DuckDBConnector("t_a", sources).run_select("select sum(units) from dealer", 1)

        assert rows == [(14,)]
        assert sources[0].profile["dropped_total_rows"] == 1
        assert sources[0].profile["row_count"] == 2

    def test_a_day_first_text_column_becomes_a_timestamp(self, tmp_path):
        sources = _csv(tmp_path, "orders.csv", "sold_on,units\n04/07/2026,10\n13/07/2026,4\n")

        assert _types(sources)["sold_on"] == "TIMESTAMP"
        assert sources[0].profile["date_range"] == {
            "sold_on": ["2026-07-04T00:00:00", "2026-07-13T00:00:00"]
        }

    def test_a_column_that_is_only_half_dates_stays_text(self, tmp_path):
        sources = _csv(
            tmp_path,
            "orders.csv",
            "sold_on,units\n04/07/2026,1\npending,2\n13/07/2026,3\nlater,4\n",
        )

        assert _types(sources)["sold_on"] == "VARCHAR"
        assert sources[0].profile["date_range"] == {}

    def test_an_iso_date_keeps_its_month(self, tmp_path):
        sources = _csv(tmp_path, "orders.csv", "month,units\n2026-07-01,10\n2026-08-01,4\n")

        assert sources[0].profile["date_range"] == {
            "month": ["2026-07-01T00:00:00", "2026-08-01T00:00:00"]
        }

    def test_month_names_stay_text(self, tmp_path):
        sources = _csv(tmp_path, "orders.csv", "month,units\nJuly,10\nAugust,4\n")

        assert _types(sources)["month"] == "VARCHAR"

    def test_an_empty_unnamed_column_is_dropped_and_a_populated_one_kept(self, tmp_path):
        (source,) = _workbook(
            tmp_path,
            "dealer.xlsx",
            {
                "July": [
                    ["dealer", None, "units", None, None, "note"],
                    ["A", "north", 10, None, None, "ok"],
                    ["B", "south", 4, None, None, "late"],
                ]
            },
        )

        assert list(source.profile["columns"]) == ["dealer", "unnamed_1", "units", "note"]


class TestNumbers:
    @pytest.mark.parametrize(
        ("values", "expected"),
        [
            (["1,23,456.50", "2,000.25"], [123456.5, 2000.25]),
            (["123,456.50", "7.25"], [123456.5, 7.25]),
            (["₹1,23,456", "Rs. 2,000", "INR 500", "rs 12"], [123456, 2000, 500, 12]),
            (["(1,234)", "500", "-20"], [-1234, 500, -20]),
            (["1,000.00 Dr", "250.00 Cr", "75.50 dr."], [-1000.0, 250.0, -75.5]),
            (["1,000.00 Dr", "250.00 Dr"], [1000.0, 250.0]),
        ],
    )
    def test_an_amount_column_becomes_numbers(self, values, expected):
        import pandas as pd

        assert _numbers(pd.Series(values)).tolist() == expected

    def test_whole_amounts_stay_whole(self):
        import pandas as pd

        assert str(_numbers(pd.Series(["1,200", "Rs. 30"])).dtype) == "Int64"

    def test_a_column_under_ninety_percent_amounts_stays_text(self):
        import pandas as pd

        column = pd.Series(["1,234", "pending", "2,000", "n/a"])

        assert _numbers(column).tolist() == ["1,234", "pending", "2,000", "n/a"]

    def test_amounts_in_a_csv_are_summed_as_numbers(self, tmp_path):
        sources = _csv(tmp_path, "ledger.csv", 'party,amount\nA,"₹1,23,456.50"\nB,"(1,000.50)"\n')

        _, rows = DuckDBConnector("t_a", sources).run_select("select sum(amount) from ledger", 1)

        assert _types(sources)["amount"] == "DOUBLE"
        assert rows == [(122456.0,)]


class TestSchema:
    def test_it_lists_the_uploaded_tables(self, dataset):
        assert dataset.list_tables() == ["q3_sales"]

    def test_it_reads_the_uploaded_columns_in_file_order(self, dataset):
        (table,) = dataset.read_tables(["q3_sales"])

        assert table.name == "q3_sales"
        assert [c.name for c in table.columns] == ["region", "product", "units", "revenue_inr"]

    def test_types_are_sniffed_once_at_ingest(self, dataset):
        (table,) = dataset.read_tables(["q3_sales"])
        types = {c.name: c.type for c in table.columns}

        assert types["units"] == "BIGINT"
        assert types["region"] == "VARCHAR"

    def test_a_table_the_file_does_not_hold_is_not_invented(self, dataset):
        assert [t.name for t in dataset.read_tables(["q3_sales", "elsewhere"])] == ["q3_sales"]

    def test_stats_are_a_size_for_the_requested_tables_only(self, tmp_path, sales):
        taken = {"q3_sales"}
        big = _csv(tmp_path, "ledger.csv", "units\n" + "1\n" * 1234, taken)
        small = _csv(tmp_path, "stock.csv", "units\n1\n", taken)
        dataset = _dataset(sales + big + small)

        assert dataset.table_stats(["ledger", "q3_sales", "elsewhere"]) == {
            "ledger": {"rows": "thousands", "rows_approx": 1200},
            "q3_sales": {"rows": "few"},
        }

    def test_the_columns_read_are_the_ones_duckdb_materialises(self, tmp_path):
        parts = [
            _csv(tmp_path, "july.csv", "units,note,only_july\n1,a,x\n")[0],
            _csv(tmp_path, "august.csv", "units,note\n2.5,3\n")[0],
        ]
        union = [FileSource(**{**vars(p), "table": "sales"}) for p in parts]

        _, materialised = DuckDBConnector("t_a", union).run_select(
            "select column_name, data_type from information_schema.columns "
            "where table_name = 'sales' order by ordinal_position",
            10,
        )
        (table,) = _dataset(union).read_tables(["sales"])

        assert [(c.name, c.type) for c in table.columns] == materialised
        assert dict(materialised) == {"units": "DOUBLE", "note": "VARCHAR", "only_july": "VARCHAR"}

    def test_row_counts_add_up_across_a_tables_parts(self, tmp_path):
        parts = [_csv(tmp_path, f"m{i}.csv", "units\n1\n2\n3\n")[0] for i in range(2)]
        union = [FileSource(**{**vars(p), "table": "sales"}) for p in parts]

        _, rows = DuckDBConnector("t_a", union).run_select("select count(*) from sales", 1)

        assert _dataset(union).table_stats(["sales"]) == {"sales": {"rows": "few"}}
        assert rows == [(6,)]


class TestOpening:
    def test_only_the_chosen_tables_are_materialised(self, tmp_path, monkeypatch):
        from app.config import get_settings

        monkeypatch.setattr(get_settings(), "file_store_dir", str(tmp_path), raising=False)
        taken: set[str] = set()
        sources = [s for i in range(20) for s in _csv(tmp_path, f"table{i}.csv", CSV, taken)]
        dataset = _dataset(sources, root=tmp_path)

        _, rows = dataset.open({"table3", "table7", "table11"}).run_select(
            "select table_name from information_schema.tables order by table_name", 50
        )

        assert rows == [("table11",), ("table3",), ("table7",)]
        assert len(dataset.list_tables()) == 20


class TestQuerying:
    def test_it_answers_a_question_about_the_file(self, connector):
        cols, rows = connector.run_select("select sum(units) as total from q3_sales", 10)

        assert cols == ["total"]
        assert rows == [(14,)]

    def test_rows_are_capped(self, connector):
        _, rows = connector.run_select("select * from q3_sales", 1)

        assert len(rows) == 1


class TestLockdown:
    """After ingest the process cannot touch the filesystem, whatever the guard thinks. This is
    the analogue of `default_transaction_read_only=on` on the Postgres connector."""

    @pytest.mark.parametrize(
        "sql",
        [
            "select * from read_csv_auto('C:/Windows/win.ini')",
            "select * from read_parquet('C:/x.parquet')",
            "select * from read_text('C:/Windows/win.ini')",
            "select * from read_blob('C:/Windows/win.ini')",
            "install httpfs",
            "copy q3_sales to 'C:/leak.csv'",
        ],
    )
    def test_the_engine_refuses_to_reach_the_filesystem(self, connector, sql):
        with pytest.raises(Exception, match=r"(?i)permission|not allowed|denied"):
            connector.run_select(sql, 10)

    def test_attaching_another_database_is_refused(self, connector, tmp_path):
        with pytest.raises(Exception, match=r"(?i)permission"):
            connector.run_select(f"attach '{(tmp_path / 'o.db').as_posix()}' as o", 10)

    def test_duckdbs_own_temp_directory_is_closed_too(self, sales, tmp_path, monkeypatch):
        import duckdb

        monkeypatch.chdir(tmp_path)
        (tmp_path / ".tmp").mkdir()
        duckdb.sql("copy (select 1 as x) to '.tmp/probe.parquet'")

        with pytest.raises(Exception, match=r"(?i)permission|not allowed|denied"):
            DuckDBConnector("t_a", sales).run_select(
                "select * from read_parquet('.tmp/probe.parquet')", 10
            )

    def test_the_lockdown_cannot_be_lifted(self, connector):
        with pytest.raises(Exception, match=r"(?i)lock"):
            connector.run_select("set enable_external_access=true", 10)

    def test_a_table_name_that_is_not_an_identifier_never_reaches_create_table(self, sales):
        bad = [
            FileSource(
                table="sales; drop table x",
                path=sales[0].path,
                file="x.csv",
                origin="upload",
                profile=sales[0].profile,
            )
        ]

        with pytest.raises(ValueError, match="unsafe table name"):
            DuckDBConnector("t_a", bad)


class TestProtocol:
    def test_it_is_a_sql_connector_that_parses_duckdb(self, connector):
        assert connector.kind == "file"
        assert connector.dialect == "duckdb"
        assert hasattr(connector, "run_select")
